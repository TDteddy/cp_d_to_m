"""
비교 결과를 Excel 파일로 생성하는 모듈
"""
import pandas as pd
from datetime import datetime
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ReportGenerator:
    """비교 결과 보고서 생성 클래스"""

    def __init__(self, output_dir: str):
        """
        Args:
            output_dir: 출력 파일 디렉토리
        """
        self.output_dir = output_dir

    def generate_excel_report(
        self,
        result_df: pd.DataFrame,
        summary: dict,
        year: int,
        month: int
    ) -> str:
        """
        비교 결과를 Excel 파일로 생성

        Args:
            result_df: 비교 결과 DataFrame
            summary: 요약 통계
            year: 연도
            month: 월

        Returns:
            str: 생성된 파일 경로
        """
        # 파일명 생성
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'판매보고서_비교결과_{year}년_{month}월_{timestamp}.xlsx'
        filepath = os.path.join(self.output_dir, filename)

        # Excel 파일 생성
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            # 요약 시트
            summary_df = pd.DataFrame([summary]).T
            summary_df.columns = ['값']
            summary_df.index.name = '항목'
            summary_df.to_excel(writer, sheet_name='요약')

            # 상세 비교 결과 시트
            result_df.to_excel(writer, sheet_name='상세 비교', index=False)

            # 차이가 있는 항목만 필터링한 시트
            diff_df = result_df[
                (result_df['매출_차이'] != 0) | (result_df['판매량_차이'] != 0)
            ]
            if len(diff_df) > 0:
                diff_df.to_excel(writer, sheet_name='차이 있는 항목', index=False)

        # 스타일 적용
        self._apply_styles(filepath)

        print(f"보고서 생성 완료: {filepath}")
        return filepath

    def _apply_styles(self, filepath: str):
        """
        Excel 파일에 스타일 적용

        Args:
            filepath: Excel 파일 경로
        """
        wb = load_workbook(filepath)

        # 모든 시트에 스타일 적용
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            # 헤더 스타일
            header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
            header_font = Font(color='FFFFFF', bold=True, size=11)
            header_alignment = Alignment(horizontal='center', vertical='center')

            # 첫 번째 행(헤더)에 스타일 적용
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_alignment

            # 차이가 있는 셀 강조 (상세 비교 시트와 차이 있는 항목 시트)
            if sheet_name in ['상세 비교', '차이 있는 항목']:
                self._highlight_differences(ws)

            # 컬럼 너비 자동 조정
            self._auto_adjust_column_width(ws)

            # 테두리 적용
            self._apply_borders(ws)

        wb.save(filepath)

    def _highlight_differences(self, ws):
        """
        차이가 있는 셀을 강조

        Args:
            ws: 워크시트
        """
        # 차이 컬럼 찾기
        headers = [cell.value for cell in ws[1]]

        diff_cols = []
        for idx, header in enumerate(headers, 1):
            if '차이' in str(header):
                diff_cols.append(idx)

        # 차이가 있는 행 강조
        red_fill = PatternFill(start_color='FFE6E6', end_color='FFE6E6', fill_type='solid')
        green_fill = PatternFill(start_color='E6FFE6', end_color='E6FFE6', fill_type='solid')

        for row in range(2, ws.max_row + 1):
            for col in diff_cols:
                cell = ws.cell(row=row, column=col)
                if cell.value is not None and cell.value != 0:
                    # 양수면 빨강, 음수면 초록
                    try:
                        value = float(cell.value)
                        if value > 0:
                            cell.fill = red_fill
                        elif value < 0:
                            cell.fill = green_fill
                    except (ValueError, TypeError):
                        pass

    def _auto_adjust_column_width(self, ws):
        """
        컬럼 너비 자동 조정

        Args:
            ws: 워크시트
        """
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)

            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass

            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    def _apply_borders(self, ws):
        """
        테두리 적용

        Args:
            ws: 워크시트
        """
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='left', vertical='center')
